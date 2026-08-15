from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import CourseSession, ScheduleAssignment, SolverRun, Teacher, TimeSlot


def solve(client: TestClient, headers: dict[str, str]) -> dict:
    response = client.post(
        "/api/v1/solver-runs",
        headers=headers,
        json={
            "wait": True,
            "time_limit_seconds": 10,
            "change_weight": 100000,
        },
    )
    assert response.status_code == 202
    return response.json()


def test_seeded_overview(client: TestClient, auth_headers: dict[str, str]) -> None:
    response = client.get("/api/v1/overview", headers=auth_headers)
    assert response.status_code == 200
    counts = response.json()["counts"]
    assert counts["teachers"] == 6
    assert counts["class_groups"] == 12
    assert counts["rooms"] == 6
    assert counts["time_slots"] == 10
    assert counts["course_sessions"] == 24


def test_admin_manages_read_only_member_accounts(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    created = client.post(
        "/api/v1/users",
        headers=auth_headers,
        json={"username": "member_demo", "password": "member-pass-2026"},
    )
    assert created.status_code == 201
    member = created.json()
    assert member["role"] == "viewer"
    assert member["is_active"] is True

    login = client.post(
        "/api/v1/auth/token",
        data={"username": "member_demo", "password": "member-pass-2026"},
    )
    assert login.status_code == 200
    member_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    assert client.get("/api/v1/teachers", headers=member_headers).status_code == 200
    assert client.get("/api/v1/users", headers=member_headers).status_code == 403
    campus = client.get("/api/v1/campuses", headers=member_headers).json()[0]
    forbidden_write = client.post(
        "/api/v1/rooms",
        headers=member_headers,
        json={
            "campus_id": campus["id"],
            "business_id": "MEMBER-FORBIDDEN-ROOM",
            "name": "成员不可新增",
            "is_active": True,
        },
    )
    assert forbidden_write.status_code == 403

    disabled = client.patch(
        f"/api/v1/users/{member['id']}/status",
        headers=auth_headers,
        json={"is_active": False},
    )
    assert disabled.status_code == 200
    assert disabled.json()["is_active"] is False
    assert client.get("/api/v1/overview", headers=member_headers).status_code == 401
    assert (
        client.post(
            "/api/v1/auth/token",
            data={"username": "member_demo", "password": "member-pass-2026"},
        ).status_code
        == 401
    )

    reenabled = client.patch(
        f"/api/v1/users/{member['id']}/status",
        headers=auth_headers,
        json={"is_active": True},
    )
    assert reenabled.status_code == 200
    reset = client.post(
        f"/api/v1/users/{member['id']}/reset-password",
        headers=auth_headers,
        json={"password": "member-new-pass-2026"},
    )
    assert reset.status_code == 204
    assert (
        client.post(
            "/api/v1/auth/token",
            data={"username": "member_demo", "password": "member-pass-2026"},
        ).status_code
        == 401
    )
    assert (
        client.post(
            "/api/v1/auth/token",
            data={"username": "member_demo", "password": "member-new-pass-2026"},
        ).status_code
        == 200
    )


def test_download_master_data_sample(client: TestClient, auth_headers: dict[str, str]) -> None:
    response = client.get("/api/v1/imports/sample.xlsx", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert {"使用说明", "课表数据源"}.issubset(workbook.sheetnames)


def test_sample_solve_publish_and_xlsx(client: TestClient, auth_headers: dict[str, str]) -> None:
    run = solve(client, auth_headers)
    assert run["model_status"] in {"OPTIMAL", "FEASIBLE"}
    schedules = client.get("/api/v1/schedules", headers=auth_headers).json()
    schedule = client.get(f"/api/v1/schedules/{schedules[0]['id']}", headers=auth_headers).json()
    assert len(schedule["assignments"]) == 24
    assert schedule["metrics"]["hard_conflicts"] == 0
    assert schedule["metrics"]["room_slot_occupancy"] >= 0

    published = client.post(f"/api/v1/schedules/{schedule['id']}/publish", headers=auth_headers)
    assert published.status_code == 200
    assert published.json()["status"] == "published"

    exported = client.get(f"/api/v1/schedules/{schedule['id']}/export.xlsx", headers=auth_headers)
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["课表", "指标"]
    assert workbook["课表"].max_row == 25


def test_publish_keeps_local_version_when_published_data_sync_partially_fails(
    client: TestClient,
    auth_headers: dict[str, str],
    monkeypatch: Any,
) -> None:
    run = solve(client, auth_headers)
    schedule = next(
        item
        for item in client.get("/api/v1/schedules", headers=auth_headers).json()
        if item["solver_run_id"] == run["id"]
    )
    calls: list[tuple[str, str]] = []

    def sync_rows(
        self: object,
        user_id: str,
        resource: str,
        rows: list[dict[str, Any]],
        workspace_id: str | None = None,
        schedule_set_id: str = "default",
    ) -> dict[str, Any]:
        del self, user_id, rows, workspace_id
        calls.append((resource, schedule_set_id))
        if resource == "public_summary":
            raise RuntimeError("公开展示汇总暂时不可写入")
        return {
            "records_read": 4,
            "records_written": 3,
            "records_created": 2,
            "records_updated": 1,
        }

    monkeypatch.setattr("app.api.FeishuService.sync_rows", sync_rows)
    published = client.post(
        f"/api/v1/schedules/{schedule['id']}/publish", headers=auth_headers
    )
    assert published.status_code == 200, published.text
    assert published.json()["status"] == "published"
    assert calls == [("schedule", "default"), ("public_summary", "default")]

    syncs = client.get("/api/v1/integrations/feishu/syncs", headers=auth_headers)
    assert syncs.status_code == 200
    automatic = [
        item
        for item in syncs.json()
        if item["detail"].get("trigger") == "version_publish"
    ]
    assert {item["resource"] for item in automatic} == {"schedule", "public_summary"}
    assert {item["status"] for item in automatic} == {"completed", "failed"}
    assert all(item["detail"]["schedule_set_id"] == "default" for item in automatic)


def test_course_session_create_edit_batch_update_and_delete(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    campus = client.get("/api/v1/campuses", headers=auth_headers).json()[0]
    class_group = client.get("/api/v1/class-groups", headers=auth_headers).json()[0]
    teacher = client.get("/api/v1/teachers", headers=auth_headers).json()[0]
    room = client.get("/api/v1/rooms", headers=auth_headers).json()[0]

    created = client.post(
        "/api/v1/course-sessions",
        headers=auth_headers,
        json={
            "campus_id": campus["id"],
            "business_id": "MANUAL-COURSE-01",
            "class_business_id": class_group["business_id"],
            "teacher_business_id": teacher["business_id"],
            "business_line": "考研",
            "product_type": "手工新增班型",
            "lesson_name": "手工新增课程",
            "session_no": 1,
            "lesson_date": "2026-10-01",
            "duration_minutes": 180,
            "fixed_start_time": "09:00",
            "fixed_end_time": "12:00",
            "original_room_business_id": room["business_id"],
        },
    )
    assert created.status_code == 201
    course = created.json()

    edited = client.put(
        f"/api/v1/course-sessions/{course['id']}",
        headers=auth_headers,
        json={"lesson_date": "2026-10-02", "original_room_business_id": room["business_id"]},
    )
    assert edited.status_code == 200
    assert edited.json()["lesson_date"] == "2026-10-02"
    assert edited.json()["teacher_business_id"] == teacher["business_id"]
    assert edited.json()["fixed_start_time"] == "09:00"

    immutable_field = client.put(
        f"/api/v1/course-sessions/{course['id']}",
        headers=auth_headers,
        json={"teacher_business_id": "OTHER-TEACHER"},
    )
    assert immutable_field.status_code == 422

    batch = client.post(
        "/api/v1/course-sessions/batch-update",
        headers=auth_headers,
        json={"object_ids": [course["id"]], "lesson_date": "2026-10-03"},
    )
    assert batch.status_code == 200
    assert batch.json()["affected_count"] == 1

    deleted = client.post(
        "/api/v1/course-sessions/batch-delete",
        headers=auth_headers,
        json={"object_ids": [course["id"]]},
    )
    assert deleted.status_code == 200
    assert deleted.json()["affected_count"] == 1


def test_infeasible_rules_return_traceable_core(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    rule_ids: list[str] = []
    for business_id, course_id in [("TEST-FIX-1", "B01-1"), ("TEST-FIX-2", "B07-1")]:
        created = client.post(
            "/api/v1/rules",
            headers=auth_headers,
            json={
                "business_id": business_id,
                "source_text": f"{course_id} 固定到 S01",
                "actor_type": "course",
                "actor_ids": [course_id],
                "constraint_type": "fixed_slot",
                "scope": {"slot_id": "S01"},
                "hardness": "hard",
            },
        )
        assert created.status_code == 201
        rule_id = created.json()["id"]
        rule_ids.append(rule_id)
        activated = client.post(
            f"/api/v1/rules/{rule_id}/transition",
            headers=auth_headers,
            json={"status": "active"},
        )
        assert activated.status_code == 200

    run = solve(client, auth_headers)
    assert run["model_status"] == "INFEASIBLE"
    assert set(run["conflict_rule_ids"]) == {"TEST-FIX-1", "TEST-FIX-2"}
    assert run["priority_rule_ids"] == ["TEST-FIX-1", "TEST-FIX-2"]
    assert run["priority_explanations"]

    for rule_id in rule_ids:
        retired = client.post(
            f"/api/v1/rules/{rule_id}/transition",
            headers=auth_headers,
            json={"status": "retired"},
        )
        assert retired.status_code == 200


def test_feishu_requires_production_configuration(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    connection = client.get("/api/v1/integrations/feishu/connection", headers=auth_headers)
    assert connection.status_code == 200
    payload = connection.json()
    assert payload["status"] == "unconfigured"
    assert payload["app_configured"] is False
    assert "应用编号" in payload["missing_fields"]
    assert payload["app_configuration"]["source"] == "none"
    assert payload["console_url"].startswith("https://open.feishu.cn/")
    sync = client.post(
        "/api/v1/integrations/feishu/sync",
        headers=auth_headers,
        json={"direction": "export", "resource": "teachers"},
    )
    assert sync.status_code == 409
    assert "填写飞书应用编号和应用密钥" in sync.json()["detail"]


def test_aily_context_proposal_confirmation_and_solve(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    aily_headers = {
        "X-Aily-Key": "test-aily-key-not-the-repo-default",
        "X-Schedule-Set-Id": "default",
    }
    context = client.get("/api/v1/aily/context", headers=aily_headers)
    assert context.status_code == 200
    assert len(context.json()["entities"]["teachers"]) == 6
    assert {item["type"] for item in context.json()["constraint_catalog"]} >= {
        "unavailable_slot",
        "consecutive_sessions",
    }

    response = client.post(
        "/api/v1/aily/rule-proposals",
        headers=aily_headers,
        json={
            "source_text": "教师甲周三不排晚课，初三冲刺班尽量连续两节；机房周四晚间检修。",
            "source_doc": "Aily 端到端基准句",
            "proposals": [
                {
                    "source_text": "教师甲周三不排晚课",
                    "actor_type": "teacher",
                    "actor_ids": ["T01"],
                    "constraint_type": "unavailable_slot",
                    "scope": {"slot_ids": ["S05", "S06"]},
                    "hardness": "hard",
                    "confidence": 0.98,
                },
                {
                    "source_text": "初三冲刺班尽量连续两节",
                    "actor_type": "class",
                    "actor_ids": ["B04", "B05", "B10"],
                    "constraint_type": "consecutive_sessions",
                    "scope": {"minimum_consecutive": 2},
                    "hardness": "soft",
                    "weight": 50,
                    "confidence": 0.88,
                },
                {
                    "source_text": "机房周四晚间检修",
                    "actor_type": "room",
                    "actor_ids": ["R03"],
                    "constraint_type": "unavailable_slot",
                    "scope": {"slot_ids": ["S07", "S08"]},
                    "hardness": "hard",
                    "confidence": 0.99,
                },
            ],
        },
    )
    assert response.status_code == 200
    proposals = response.json()
    assert len(proposals) == 3
    assert {item["status"] for item in proposals} == {"awaiting_confirmation"}

    for proposal in proposals:
        transition = client.post(
            f"/api/v1/rules/{proposal['id']}/transition",
            headers=auth_headers,
            json={"status": "active", "reason": "基准句人工确认"},
        )
        assert transition.status_code == 200

    run = solve(client, auth_headers)
    assert run["model_status"] in {"OPTIMAL", "FEASIBLE"}

    for proposal in proposals:
        retired = client.post(
            f"/api/v1/rules/{proposal['id']}/transition",
            headers=auth_headers,
            json={"status": "retired", "reason": "测试清理"},
        )
        assert retired.status_code == 200


def test_product_loop_calendar_assistant_export_and_public_summary(
    client: TestClient,
    auth_headers: dict[str, str],
    monkeypatch: Any,
) -> None:
    run = solve(client, auth_headers)
    schedules = client.get("/api/v1/schedules", headers=auth_headers).json()
    summary = next(item for item in schedules if item["solver_run_id"] == run["id"])
    schedule = client.get(f"/api/v1/schedules/{summary['id']}", headers=auth_headers).json()
    assert summary["assignment_count"] == len(schedule["assignments"])
    published = client.post(f"/api/v1/schedules/{schedule['id']}/publish", headers=auth_headers)
    assert published.status_code == 200

    with SessionLocal() as db:
        assignment = db.scalar(
            select(ScheduleAssignment).where(
                ScheduleAssignment.schedule_version_id == schedule["id"]
            )
        )
        assert assignment is not None
        course = db.get(CourseSession, assignment.course_session_id)
        assert course is not None
        teacher = db.scalar(
            select(Teacher).where(Teacher.business_id == course.teacher_business_id)
        )
        assert teacher is not None
        assignment.lesson_date = date(2026, 9, 5)
        db.add(
            TimeSlot(
                campus_id=course.campus_id,
                business_id="S-TEST-0900",
                weekday="周六",
                start_time="09:00",
                end_time="12:00",
                kind="测试实际时段",
                sequence=999,
            )
        )
        assignment.slot_business_id = "S-TEST-0900"
        course.business_line = "内部业务线"
        course.product_type = "内部产品班型"
        course.fixed_start_time = "09:00"
        course.fixed_end_time = "12:00"
        course.original_room_business_id = assignment.room_business_id
        course.calendar_user_id = None
        teacher.calendar_user_id = "ou_calendar_teacher"
        db.commit()

    exported = client.get(f"/api/v1/schedules/{schedule['id']}/export.xlsx", headers=auth_headers)
    workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
    headers = [cell.value for cell in next(workbook["课表"].iter_rows(max_row=1))]
    assert {
        "上课日期",
        "业务线",
        "产品班型",
        "固定开始时间",
        "固定结束时间",
        "原始教室ID",
        "最终教室ID",
        "具体日程账号",
    } <= set(headers)

    context = client.get(
        "/api/v1/aily/context",
        headers={
            "X-Aily-Key": "test-aily-key-not-the-repo-default",
            "X-Schedule-Set-Id": "default",
        },
    )
    mapped = [
        item
        for item in context.json()["entities"]["courses"]
        if item["calendar_mapping_status"] == "mapped"
    ]
    assert mapped and mapped[0]["fixed_start_time"]

    monkeypatch.setattr("app.api.enqueue_solver_run", lambda run_id: None)
    monkeypatch.setattr("app.api.settings.aily_app_id", "spring_test")
    monkeypatch.setattr("app.api.settings.aily_skill_id", "skill_test")
    monkeypatch.setattr(
        "app.api.FeishuService.start_aily_skill",
        lambda *args, **kwargs: {
            "business_lines": ["内部业务线"],
            "product_types": [],
            "class_business_ids": [],
            "date_from": None,
            "date_to": None,
            "date_window_days": 2,
            "recognized_rules": [
                "固定时段不可调整",
                "同一教室真实时间区间不可重叠",
                "优先最小化日期和教室变更",
            ],
        },
    )
    interpreted = client.post(
        "/api/v1/assistant/interpret",
        headers=auth_headers,
        json={"instruction": "内部业务线在固定时段不变的前提下尽量不变，日期范围 2 天"},
    )
    assert interpreted.status_code == 200
    assert interpreted.json()["source"] == "feishu_aily"
    assert interpreted.json()["business_lines"] == ["内部业务线"]
    assert interpreted.json()["date_window_days"] == 2
    assert "固定时段不可调整" in interpreted.json()["recognized_rules"]
    assert "fixed_time" in interpreted.json()["solver_rules"]
    assistant = client.post(
        "/api/v1/assistant/solve",
        headers=auth_headers,
        json={
            "instruction": "仅重排内部业务线，日期前后允许移动两天",
            "business_lines": ["内部业务线"],
            "date_window_days": 2,
        },
    )
    assert assistant.status_code == 202
    with SessionLocal() as db:
        assistant_run = db.get(SolverRun, assistant.json()["id"])
        assert assistant_run is not None
        assert assistant_run.request_payload["instruction"].startswith("仅重排")
        assert assistant_run.request_payload["assistant_entry"] is True
        assert "room_no_overlap" in assistant_run.request_payload["solver_rules"]
        assert assistant_run.request_payload["parent_schedule_id"] == schedule["id"]
        assert len(assistant_run.request_payload["previous_assignments"]) == len(
            schedule["assignments"]
        )

    def freebusy(*args: Any, **kwargs: Any) -> dict[str, Any]:
        assert len(kwargs["user_ids"]) <= 10
        return {
            "freebusy_lists": [
                {
                    "user_id": "ou_calendar_teacher",
                    "freebusy_list": [
                        {
                            "start_time": "2026-09-05T10:00:00+08:00",
                            "end_time": "2026-09-05T11:00:00+08:00",
                        }
                    ],
                }
            ]
        }

    monkeypatch.setattr("app.api.FeishuService.batch_freebusy", freebusy)
    monkeypatch.setattr(
        "app.api.FeishuService.create_calendar_event",
        lambda *args, **kwargs: {"event": {"event_id": "event-test-1"}},
    )
    monkeypatch.setattr(
        "app.api.FeishuService.add_event_attendee", lambda *args, **kwargs: {"attendees": []}
    )
    dry_run = client.post(
        f"/api/v1/schedules/{schedule['id']}/calendar-publish",
        headers=auth_headers,
        json={"dry_run": True},
    )
    assert dry_run.status_code == 200
    assert dry_run.json()["published"] == 0
    assert dry_run.json()["would_publish"] == 1
    assert dry_run.json()["conflict_count"] == 1

    created = client.post(
        f"/api/v1/schedules/{schedule['id']}/calendar-publish",
        headers=auth_headers,
        json={},
    )
    assert created.status_code == 200
    assert created.json()["published"] == 1
    assert created.json()["conflict_count"] == 1
    bindings = client.get(
        f"/api/v1/schedules/{schedule['id']}/calendar-bindings", headers=auth_headers
    ).json()
    assert bindings[0]["status"] == "published_with_conflict"
    repeated = client.post(
        f"/api/v1/schedules/{schedule['id']}/calendar-publish",
        headers=auth_headers,
        json={},
    )
    assert repeated.json()["published"] == 0
    assert repeated.json()["existing"] == 1

    from app.api import export_resource_rows

    with SessionLocal() as db:
        public_rows = export_resource_rows(db, "public_summary")
    assert any(item["指标名称"] == "总课次" for item in public_rows)
    assert "内部业务线" not in str(public_rows)
    assert "内部产品班型" not in str(public_rows)


def test_assistant_rejects_unknown_scope(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        "/api/v1/assistant/solve",
        headers=auth_headers,
        json={
            "instruction": "仅重排不存在的业务线",
            "business_lines": ["不存在的业务线"],
        },
    )
    assert response.status_code == 422
    assert "未知业务实体" in response.text


def test_rule_status_cannot_be_forced_to_active_on_create(client, auth_headers) -> None:
    """状态只能由 transition 推进：创建接口收到 status 字段应当直接拒绝。"""
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "越权直接生效",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "declared_constraint",
            "scope": {},
            "hardness": "hard",
            "status": "active",
        },
    )
    assert response.status_code == 422


def test_aily_endpoints_reject_a_wrong_key(client) -> None:
    response = client.get("/api/v1/aily/context", headers={"X-Aily-Key": "wrong-key"})
    assert response.status_code == 401


def test_audit_logs_are_admin_only(client, auth_headers) -> None:
    created = client.post(
        "/api/v1/users",
        headers=auth_headers,
        json={"username": "audit_viewer", "password": "viewer-password-1"},
    )
    assert created.status_code == 201
    assert created.json()["role"] == "viewer"
    token = client.post(
        "/api/v1/auth/token",
        data={"username": "audit_viewer", "password": "viewer-password-1"},
    ).json()["access_token"]

    response = client.get("/api/v1/audit-logs", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 403


def test_admin_can_assign_roles(client, auth_headers) -> None:
    created = client.post(
        "/api/v1/users",
        headers=auth_headers,
        json={"username": "new_scheduler", "password": "scheduler-password-1", "role": "scheduler"},
    )
    assert created.status_code == 201
    assert created.json()["role"] == "scheduler"

    changed = client.patch(
        f"/api/v1/users/{created.json()['id']}/role",
        headers=auth_headers,
        json={"role": "approver"},
    )
    assert changed.status_code == 200
    assert changed.json()["role"] == "approver"


def test_password_change_revokes_previously_issued_tokens(client, auth_headers) -> None:
    client.post(
        "/api/v1/users",
        headers=auth_headers,
        json={"username": "rotating_user", "password": "old-password-123"},
    )
    old_token = client.post(
        "/api/v1/auth/token",
        data={"username": "rotating_user", "password": "old-password-123"},
    ).json()["access_token"]
    old_headers = {"Authorization": f"Bearer {old_token}"}
    assert client.get("/api/v1/auth/me", headers=old_headers).status_code == 200

    changed = client.post(
        "/api/v1/auth/change-password",
        headers=old_headers,
        json={"current_password": "old-password-123", "new_password": "brand-new-password-1"},
    )
    assert changed.status_code == 204
    assert client.get("/api/v1/auth/me", headers=old_headers).status_code == 401
    assert (
        client.post(
            "/api/v1/auth/token",
            data={"username": "rotating_user", "password": "brand-new-password-1"},
        ).status_code
        == 200
    )


def test_repeated_login_failures_lock_the_account(client, auth_headers) -> None:
    client.post(
        "/api/v1/users",
        headers=auth_headers,
        json={"username": "brute_target", "password": "correct-password-1"},
    )
    for _ in range(8):
        assert (
            client.post(
                "/api/v1/auth/token",
                data={"username": "brute_target", "password": "wrong"},
            ).status_code
            == 401
        )
    locked = client.post(
        "/api/v1/auth/token",
        data={"username": "brute_target", "password": "correct-password-1"},
    )
    assert locked.status_code == 429


def test_constraint_catalog_covers_the_types_the_solver_implements(client, auth_headers) -> None:
    """目录曾经漏掉求解器已实现的日期/教室类型，导致这些规则根本建不出来。"""
    response = client.get("/api/v1/rules/constraint-catalog", headers=auth_headers)
    assert response.status_code == 200
    catalog = {item["type"]: item for item in response.json()}

    for constraint_type in (
        "fixed_date",
        "preferred_date",
        "date_range",
        "date_window",
        "allowed_date_range",
        "preferred_room",
        "forbidden_room",
        "unavailable_room",
    ):
        assert constraint_type in catalog, f"{constraint_type} 求解器已实现，目录必须收录"
        assert catalog[constraint_type]["solver_paths"], f"{constraint_type} 必须声明生效路径"

    # 只登记留痕的类型要如实标注为不进入模型，前端靠这个字段提示教务。
    assert catalog["declared_constraint"]["solver_paths"] == {"hard": [], "soft": []}
    # 声明的生效路径不能超出该类型允许的硬软属性。
    for entry in catalog.values():
        assert set(entry["solver_paths"]) <= set(entry["hardness"])


def test_date_rules_can_now_be_created_through_the_api(client, auth_headers) -> None:
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "business_id": "TEST-FIXED-DATE",
            "source_text": "B01-1 必须排在 2026-09-09",
            "actor_type": "course",
            "actor_ids": ["B01-1"],
            "constraint_type": "fixed_date",
            "scope": {"date": "2026-09-09"},
            "hardness": "hard",
        },
    )
    assert response.status_code == 201
    assert response.json()["status"] == "awaiting_confirmation"


def test_rule_scope_rejects_a_malformed_date(client, auth_headers) -> None:
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "日期写错了",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "fixed_date",
            "scope": {"date": "2026-13-45"},
            "hardness": "hard",
        },
    )
    assert response.status_code == 422
    assert "不是合法日期" in response.text


def test_rule_scope_rejects_an_inverted_date_range(client, auth_headers) -> None:
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "区间反了",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "date_range",
            "scope": {"date_from": "2026-09-20", "date_to": "2026-09-10"},
            "hardness": "hard",
        },
    )
    assert response.status_code == 422
    assert "date_from 必须早于或等于 date_to" in response.text


def test_rule_scope_rejects_a_negative_day_count(client, auth_headers) -> None:
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "浮动天数为负",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "date_range",
            "scope": {"date_window_days": -3},
            "hardness": "hard",
        },
    )
    assert response.status_code == 422


def test_global_rule_scope_entities_are_validated(client, auth_headers) -> None:
    """actor_ids 为空的全局规则以前会跳过 scope 实体校验，让不存在的时段进入求解。"""
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "全局禁排一个不存在的时段",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "forbidden_slot",
            "scope": {"slot_ids": ["S-DOES-NOT-EXIST"]},
            "hardness": "hard",
        },
    )
    assert response.status_code == 422
    assert "不存在的时段" in response.text


def test_global_rule_scope_rejects_a_missing_room(client, auth_headers) -> None:
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "全局禁用一个不存在的教室",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "forbidden_room",
            "scope": {"room_ids": ["R-DOES-NOT-EXIST"]},
            "hardness": "hard",
        },
    )
    assert response.status_code == 422
    assert "不存在的教室" in response.text


def test_rule_rejects_a_hardness_the_type_does_not_support(client, auth_headers) -> None:
    response = client.post(
        "/api/v1/rules",
        headers=auth_headers,
        json={
            "source_text": "偏好日期只能是软约束",
            "actor_type": "system",
            "actor_ids": [],
            "constraint_type": "preferred_date",
            "scope": {"date": "2026-09-09"},
            "hardness": "hard",
        },
    )
    assert response.status_code == 422


def test_infeasible_run_explanation_translates_conflicts_into_business_wording(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    """无解时的解释必须把 SYSTEM-* 和教务规则都翻成人话，且不依赖 AI 可用。"""
    rule_ids: list[str] = []
    for business_id, course_id in [("TEST-EXPLAIN-1", "B01-1"), ("TEST-EXPLAIN-2", "B07-1")]:
        created = client.post(
            "/api/v1/rules",
            headers=auth_headers,
            json={
                "business_id": business_id,
                "source_text": f"{course_id} 必须固定在周一晚间第一节",
                "actor_type": "course",
                "actor_ids": [course_id],
                "constraint_type": "fixed_slot",
                "scope": {"slot_id": "S01"},
                "hardness": "hard",
            },
        )
        assert created.status_code == 201
        rule_id = created.json()["id"]
        rule_ids.append(rule_id)
        assert (
            client.post(
                f"/api/v1/rules/{rule_id}/transition",
                headers=auth_headers,
                json={"status": "active"},
            ).status_code
            == 200
        )

    run = solve(client, auth_headers)
    assert run["model_status"] == "INFEASIBLE"

    explained = client.post(f"/api/v1/solver-runs/{run['id']}/explanation", headers=auth_headers)
    assert explained.status_code == 200, explained.text
    payload = explained.json()
    joined = "\n".join(payload["explanation"])
    # 教务规则用它自己的原文说话，而不是回显规则标识。
    assert "必须固定在周一晚间第一节" in joined
    assert payload["next_actions"]

    # 解释随求解任务一起返回，前端轮询同一个接口就能拿到。
    fetched = client.get(f"/api/v1/solver-runs/{run['id']}", headers=auth_headers)
    assert fetched.status_code == 200
    assert fetched.json()["explanation"]["headline"] == payload["headline"]

    for rule_id in rule_ids:
        assert (
            client.post(
                f"/api/v1/rules/{rule_id}/transition",
                headers=auth_headers,
                json={"status": "retired"},
            ).status_code
            == 200
        )
