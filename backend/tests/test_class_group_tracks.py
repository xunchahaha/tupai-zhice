"""班级的多班型 / 多教师。

走班制下一个班级里同时存在若干条轨道（选数学的和不选的分流）。以前导入器用众数把
班型、业务线、教师各自压成一个单值，压出来的三个值彼此不保证来自同一批课次——郑州
的暑期集训营OMO4班就被写成「班型=无数学 + 教师=数学教研组」这种自相矛盾的行。
这里守住的是：接口必须把全部班型和全部教师都吐出来，并且能说清哪门课谁教。
"""

from __future__ import annotations

import os
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from app.api import class_group_responses
from app.db import Base
from app.models import ClassGroup
from app.services.converter_zhengzhou import SHEET_NAME, _read_rows, import_schedule_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
# 真实课表数据不进仓库：默认找 data/private/（已 gitignore），也可用环境变量 TUPAI_OFFICIAL_WORKBOOK
# 指向本地文件；文件不存在时依赖它的用例自动跳过。
OFFICIAL_WORKBOOK = Path(
    os.environ.get("TUPAI_OFFICIAL_WORKBOOK")
    or PROJECT_ROOT / "data" / "private" / "郑州考研公职专升本课表数据源_教室班级标签版.xlsx"
)
OMO4 = "暑期集训营OMO4班"

HEADERS = (
    "标准业务线",
    "标准产品班型",
    "集训营班级标签",
    "教室标签",
    "课表编排来源",
    "编排阶段",
    "计划课次",
    "计划课时",
    "课次序号",
    "课节名称",
    "上课日期",
    "上课时段",
    "课节时长(小时)",
    "授课教师",
)


@pytest.fixture
def db() -> Iterator[Session]:
    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(bind=engine)
    factory = sessionmaker(bind=engine, autoflush=False, future=True)
    with factory() as session:
        yield session
    engine.dispose()


def _create_course(
    client: TestClient,
    headers: dict[str, str],
    campus_id: str,
    **overrides: Any,
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "campus_id": campus_id,
        "business_line": "考研",
        "session_no": 1,
        "duration_minutes": 180,
    }
    payload.update(overrides)
    response = client.post("/api/v1/course-sessions", headers=headers, json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def test_class_group_reports_every_product_type_and_teacher(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    """一个班三条轨道：含数学的数学课、无数学的政治课和英语课。

    这是 OMO4班 的真实结构。旧实现只会显示其中一个班型和一个教师。
    """
    campus_id = client.get("/api/v1/campuses", headers=auth_headers).json()[0]["id"]
    created = client.post(
        "/api/v1/class-groups",
        headers=auth_headers,
        json={"campus_id": campus_id, "business_id": "TRACK-OMO4", "name": "轨道测试OMO4班"},
    )
    assert created.status_code == 201, created.text
    class_group = created.json()

    tracks = [
        ("考研·暑期强化（含数学）", "数学", "轨道测试数学教研组", 2),
        ("考研·暑期强化（无数学）", "政治", "轨道测试政治教研组", 3),
        ("考研·暑期强化（无数学）", "英语", "轨道测试英语教研组", 1),
    ]
    course_ids: list[str] = []
    for product_type, subject, teacher, count in tracks:
        for index in range(count):
            course_ids.append(
                _create_course(
                    client,
                    auth_headers,
                    campus_id,
                    business_id=f"TRACK-{subject}-{index}",
                    class_business_id="TRACK-OMO4",
                    teacher_business_id=teacher,
                    product_type=product_type,
                    subject=subject,
                    lesson_name=f"{subject}·轨道测试课节{index}",
                )["id"]
            )

    rows = {
        item["id"]: item
        for item in client.get("/api/v1/class-groups", headers=auth_headers).json()
    }
    row = rows[class_group["id"]]

    assert row["product_types"] == [
        "考研·暑期强化（含数学）",
        "考研·暑期强化（无数学）",
    ]
    assert row["teacher_business_ids"] == [
        "轨道测试政治教研组",
        "轨道测试数学教研组",
        "轨道测试英语教研组",
    ]
    assert row["business_lines"] == ["考研"]
    assert row["subjects"] == ["政治", "数学", "英语"]
    assert row["session_count"] == 6
    # 两个平铺数组说不清「哪门课谁教」，轨道才说得清。
    assert {
        (item["product_type"], item["subject"], item["teacher_business_id"])
        for item in row["tracks"]
    } == {(product_type, subject, teacher) for product_type, subject, teacher, _ in tracks}
    assert sum(item["session_count"] for item in row["tracks"]) == 6

    cleanup = client.post(
        "/api/v1/course-sessions/batch-delete",
        headers=auth_headers,
        json={"object_ids": course_ids},
    )
    assert cleanup.status_code == 200, cleanup.text
    removed = client.delete(
        f"/api/v1/master-data/class-groups/{class_group['id']}", headers=auth_headers
    )
    assert removed.status_code == 204, removed.text


def test_class_group_write_payload_no_longer_accepts_guessed_single_values(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    """班型/业务线/教师不再是班级的可写字段，写进来就该被挡住。

    留一个「看起来能填」的入口，等于允许有人手填一个和课次矛盾的值。
    """
    campus_id = client.get("/api/v1/campuses", headers=auth_headers).json()[0]["id"]
    response = client.post(
        "/api/v1/class-groups",
        headers=auth_headers,
        json={
            "campus_id": campus_id,
            "business_id": "TRACK-REJECT",
            "name": "不该被创建的班",
            "grade": "手填班型",
            "teacher_business_id": "手填教师",
        },
    )
    assert response.status_code == 422, response.text
    rows = client.get("/api/v1/class-groups", headers=auth_headers).json()
    assert all(item["business_id"] != "TRACK-REJECT" for item in rows)
    assert all("grade" not in item and "teacher_business_id" not in item for item in rows)


def test_official_omo4_class_shows_both_product_types(db: Session, tmp_path: Path) -> None:
    """用郑州源表里 OMO4班 的真实行跑一遍导入，含数学和无数学必须同时出现。"""
    if not OFFICIAL_WORKBOOK.exists():
        pytest.skip("郑州官方数据文件不在工作区")

    workbook = load_workbook(OFFICIAL_WORKBOOK, data_only=True, read_only=True)
    sheet = workbook[SHEET_NAME]
    rows = [
        row
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if str(row[2] or "").strip() == OMO4
    ]
    workbook.close()
    assert rows, "源表里应当有 OMO4班 的行"

    path = tmp_path / "omo4.xlsx"
    target = Workbook()
    target_sheet = target.active
    target_sheet.title = SHEET_NAME
    target_sheet.append(HEADERS)
    for row in rows:
        target_sheet.append(row)
    target.save(path)

    import_schedule_workbook(db, path)
    db.commit()

    classes = list(db.scalars(select(ClassGroup)))
    assert [item.business_id for item in classes] == [OMO4]
    response = class_group_responses(db, classes)[0]

    assert response.product_types == [
        "考研·暑期强化（含数学）",
        "考研·暑期强化（无数学）",
    ]
    assert response.teacher_business_ids == [
        "郑州考研政治教研组",
        "郑州考研数学教研组",
        "郑州考研英语教研组",
    ]
    # 数学课只挂在含数学班型下，政治/英语只挂在无数学下——这才是走班分流。
    by_subject = {item.subject: item for item in response.tracks}
    assert by_subject["数学"].product_type == "考研·暑期强化（含数学）"
    assert by_subject["政治"].product_type == "考研·暑期强化（无数学）"
    assert by_subject["英语"].product_type == "考研·暑期强化（无数学）"
    assert response.session_count == sum(item.session_count for item in response.tracks)


def test_official_workbook_class_row_no_longer_contradicts_itself() -> None:
    """回归旧口径：众数会给 OMO4班 配出「班型无数学 + 教师数学教研组」。

    源数据里这个矛盾组合仍在，所以这个用例守的是「不许再用众数猜」这件事本身。
    """
    if not OFFICIAL_WORKBOOK.exists():
        pytest.skip("郑州官方数据文件不在工作区")

    source_rows, _ = _read_rows(OFFICIAL_WORKBOOK)
    omo4_rows = [row for row in source_rows if row["班级标签"] == OMO4]
    product_types = {row["产品班型"] for row in omo4_rows}
    teachers = {row["授课教师"] for row in omo4_rows}

    assert len(product_types) == 2, product_types
    assert len(teachers) == 3, teachers
    # 行数最多的班型是「无数学」，行数最多的教师是「数学教研组」——两个众数各自算出来，
    # 拼在一起就是一行假数据。
    no_math = sum(1 for row in omo4_rows if row["产品班型"].endswith("（无数学）"))
    math_group = sum(1 for row in omo4_rows if row["授课教师"] == "郑州考研数学教研组")
    assert no_math > len(omo4_rows) - no_math
    assert math_group == max(
        sum(1 for row in omo4_rows if row["授课教师"] == name) for name in teachers
    )
