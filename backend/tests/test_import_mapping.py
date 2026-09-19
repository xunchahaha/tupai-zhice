from __future__ import annotations

import json
import uuid
from collections.abc import Iterator
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.services.converter_core import parse_template_records
from app.services.import_mapping import (
    CONFIDENCE_THRESHOLD,
    ColumnMapping,
    build_records,
    column_samples,
    detect_header_candidates,
    parse_uploaded_workbook,
    pick_default_sheet,
    suggest_mapping,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SAMPLE_WORKBOOK = PROJECT_ROOT / "data" / "imports" / "sample.xlsx"

HEADERS = (
    "标准业务线",
    "标准产品班型",
    "集训营班级标签",
    "教室标签",
    "课表编排来源",
    "编排阶段",
    "计划课次",
    "计划课时",
    "课次序号",
    "课节名称",
    "上课日期",
    "上课时段",
    "课节时长(小时)",
    "授课教师",
)

# 乱序重命名：模拟教务系统导出的真实表格——列序打乱、表头是业务同义词。
RENAME = {
    "标准业务线": "业务线",
    "标准产品班型": "班型",
    "集训营班级标签": "班级名称",
    "教室标签": "上课教室",
    "课表编排来源": "来源",
    "编排阶段": "阶段",
    "计划课次": "课次数量",
    "计划课时": "总课时",
    "课次序号": "讲次",
    "课节名称": "课节主题",
    "上课日期": "日期",
    "上课时段": "时段",
    # 全角括号：只应被「规范化」层命中，别名表不收录全角变体。
    "课节时长(小时)": "课节时长（小时）",
    "授课教师": "任课教师",
}
COLUMN_ORDER = (11, 8, 13, 0, 6, 2, 12, 9, 1, 5, 10, 3, 7, 4)


def _sample_rows() -> list[tuple[Any, ...]]:
    """读取官方模板 sample.xlsx 的真实数据行作为基准数据。"""
    workbook = load_workbook(SAMPLE_WORKBOOK, data_only=True)
    sheet = workbook["课表数据源"]
    rows = [tuple(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
    workbook.close()
    return rows


def _renamed_grid() -> list[list[Any]]:
    """标题行 + 乱序重命名表头 + 真实数据行 + 缺日期的坏行 + 备注列。"""
    rows = [list(row) for row in _sample_rows()]
    broken = list(rows[0])
    broken[10] = None  # 缺上课日期 → 行级校验必须报出来
    grid: list[list[Any]] = [
        ["某机构 2026 年春季课表导出"],
        [RENAME[HEADERS[index]] for index in COLUMN_ORDER] + ["备注"],
    ]
    for row in [*rows, broken]:
        grid.append([row[index] for index in COLUMN_ORDER] + ["手工备注"])
    return grid


def _grid_to_xlsx(grid: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    for row in grid:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _samples_for(grid: list[list[Any]], header_row_index: int) -> list[list[Any]]:
    return column_samples(grid, header_row_index)


# ---------------------------------------------------------------------------
# 四层匹配：单元测试
# ---------------------------------------------------------------------------


def _mapping_by_column(mapping: list[ColumnMapping]) -> dict[str, ColumnMapping]:
    return {item.column: item for item in mapping}


def test_exact_alias_normalized_layers_hit_in_order() -> None:
    grid = [
        ["任课教师", "上 课 日 期", "课节时长（小时）", "备注"],
        ["王老师", "2026-03-08", "3", "x"],
        ["数学教研组", "2026-03-09", "3", "y"],
    ]
    mapping, ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    by_column = _mapping_by_column(mapping)

    assert ai_used is False
    assert by_column["任课教师"].target == "授课教师"
    assert by_column["任课教师"].matched_by == "alias"
    assert by_column["任课教师"].confidence == 0.95
    assert by_column["上 课 日 期"].target == "上课日期"
    assert by_column["上 课 日 期"].matched_by == "normalized"
    assert by_column["课节时长（小时）"].target == "课节时长(小时)"
    assert by_column["课节时长（小时）"].confidence == 0.9
    assert by_column["备注"].target is None
    assert by_column["备注"].matched_by == "unmatched"


def test_exact_layer_hits_template_headers_verbatim() -> None:
    grid = [
        list(HEADERS[:4]),
        ["考研", "考研·全年集训", "全年集训营一班", "教室-510"],
        ["考研", "考研·全年集训", "全年集训营一班", "教室-510"],
    ]
    mapping, _ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    by_column = _mapping_by_column(mapping)

    assert by_column["标准业务线"].matched_by == "exact"
    assert by_column["标准业务线"].confidence == 1.0
    assert by_column["标准产品班型"].target == "标准产品班型"
    assert by_column["集训营班级标签"].target == "集训营班级标签"
    assert by_column["教室标签"].target == "教室标签"


def test_one_field_only_maps_to_one_column() -> None:
    grid = [
        ["讲次", "序号", "课节名称"],
        [1, 1, "数学·基础"],
        [2, 2, "数学·强化"],
    ]
    mapping, _ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    by_column = _mapping_by_column(mapping)

    assert by_column["讲次"].target == "课次序号"
    assert by_column["序号"].target is None
    assert "占用" in by_column["序号"].rationale


def test_fuzzy_layer_suggests_with_decayed_confidence() -> None:
    grid = [
        ["授课教师姓名"],
        ["王老师"],
        ["李老师"],
    ]
    mapping, _ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    item = mapping[0]

    assert item.target == "授课教师"
    assert item.matched_by == "fuzzy"
    assert CONFIDENCE_THRESHOLD <= item.confidence < 0.8


def test_shape_check_demotes_columns_whose_data_does_not_fit() -> None:
    # 表头完全像「上课时段」，但数据是上午/下午——形状校验必须拦下这种自信的错配。
    grid = [
        ["上课时段", "上课日期"],
        ["上午", "2026-03-08"],
        ["下午", "2026-03-09"],
        ["晚上", "2026-03-10"],
    ]
    mapping, _ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    by_column = _mapping_by_column(mapping)

    assert by_column["上课时段"].target is None
    assert by_column["上课时段"].confidence < CONFIDENCE_THRESHOLD
    assert "形状校验不通过" in by_column["上课时段"].rationale
    assert by_column["上课日期"].target == "上课日期"


def test_low_confidence_columns_are_never_guessed() -> None:
    grid = [
        ["其他", "备注信息"],
        ["a", "b"],
        ["c", "d"],
    ]
    mapping, _ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    assert all(item.target is None for item in mapping)
    assert all(item.matched_by == "unmatched" for item in mapping)


def test_header_candidates_prefer_known_field_rows_over_titles() -> None:
    grid = _renamed_grid()
    candidates = detect_header_candidates(grid)

    assert candidates, "至少应给出一个候选"
    assert candidates[0].row_index == 1
    assert candidates[0].sample[0] == "时段"


def test_parse_uploaded_workbook_supports_xlsx_and_csv() -> None:
    grid = _renamed_grid()
    grids = parse_uploaded_workbook("renamed.xlsx", _grid_to_xlsx(grid))
    assert list(grids) == ["Sheet"]
    # openpyxl 会把短行（标题行）用 None 补齐到最宽行，只比对完整表头行与行数。
    assert grids["Sheet"][1] == grid[1]
    assert len(grids["Sheet"]) == len(grid)

    csv_lines = [",".join(str(cell) if cell is not None else "" for cell in row) for row in grid]
    payload = ("﻿" + "\n".join(csv_lines)).encode("utf-8")
    csv_grids = parse_uploaded_workbook("renamed.csv", payload)
    assert list(csv_grids) == ["csv"]
    assert csv_grids["csv"][1][0] == grid[1][0]

    gbk_payload = "班级,日期\n三年二班,2026-03-08\n".encode("gb18030")
    assert parse_uploaded_workbook("export.csv", gbk_payload)["csv"][1][1] == "2026-03-08"


def test_pick_default_sheet_prefers_the_data_sheet() -> None:
    grids = {
        "说明": [["使用说明"], ["第 1 行", "第 2 行"]],
        "数据": [["a", "b", "c"], ["1", "2", "3"], ["4", "5", "6"]],
    }
    assert pick_default_sheet(grids) == "数据"


def test_build_records_and_row_level_issues_reuse_the_core_parser() -> None:
    grid = _renamed_grid()
    # 网格第 j 列来自原始第 COLUMN_ORDER[j] 列，目标即该列的规范字段名。
    targets = {index: HEADERS[source] for index, source in enumerate(COLUMN_ORDER)}
    records, row_numbers = build_records(grid, 1, targets)
    assert [record["授课教师"] for record in records[:2]] == [
        str(row[13]) for row in _sample_rows()[:2]
    ]
    parsed, skipped, blank_rows = parse_template_records(records, row_numbers, max_skipped=200)

    assert blank_rows == 0
    assert len(parsed) == len(_sample_rows())
    # 坏行在最后一行：行号 = 标题(1) + 表头(1) + 数据行数 + 1。
    assert skipped == [
        {"行号": 2 + len(_sample_rows()) + 1, "原因": "缺少上课日期或上课时段"}
    ]


def test_sample_shape_accepts_datetime_strings_with_time_suffix() -> None:
    # CSV 里日期常带 00:00:00 时间尾巴，形状校验不能因此误杀。
    grid = [
        ["上课日期"],
        [str(datetime(2026, 3, 8, 0, 0, 0))],
        ["2026-03-09"],
    ]
    mapping, _ai_used = suggest_mapping(grid[0], _samples_for(grid, 0))
    assert mapping[0].target == "上课日期"
    assert mapping[0].confidence >= CONFIDENCE_THRESHOLD


# ---------------------------------------------------------------------------
# preview / commit API
# ---------------------------------------------------------------------------


@pytest.fixture(name="import_set_headers")
def import_set_headers_fixture(
    client: TestClient, auth_headers: dict[str, str]
) -> Iterator[dict[str, str]]:
    """每个用例一套独立课表方案，避免污染共享测试库的其他用例。"""
    created = client.post(
        "/api/v1/schedule-sets",
        headers=auth_headers,
        json={"name": f"智能导入-{uuid.uuid4().hex[:8]}"},
    )
    assert created.status_code in (200, 201), created.text
    headers = {**auth_headers, "X-Schedule-Set-Id": created.json()["id"]}
    yield headers


def _preview(client: TestClient, headers: dict[str, str], payload: bytes) -> dict[str, Any]:
    response = client.post(
        "/api/v1/imports/preview",
        headers=headers,
        files={
            "file": (
                "renamed.xlsx",
                payload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def _mapping_input(preview: dict[str, Any], *, skip_note: bool = True) -> str:
    columns = [
        {"column": item["column"], "column_index": item["column_index"], "target": item["target"]}
        for item in preview["mapping"]
        if item["target"] is not None
    ]
    if not skip_note:
        columns += [
            {"column": item["column"], "column_index": item["column_index"], "target": None}
            for item in preview["mapping"]
            if item["target"] is None
        ]
    return json.dumps(
        {
            "sheet": preview["selected_sheet"],
            "header_row_index": preview["header_row_index"],
            "columns": columns,
        },
        ensure_ascii=False,
    )


def test_preview_maps_renamed_and_shuffled_headers(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    preview = _preview(client, auth_headers, _grid_to_xlsx(_renamed_grid()))

    assert preview["selected_sheet"] == "Sheet"
    assert preview["header_row_index"] == 1
    assert [item["row_index"] for item in preview["header_candidates"]][0] == 1

    targets = {item["target"] for item in preview["mapping"] if item["target"] is not None}
    assert targets == set(HEADERS)
    assert all(
        item["confidence"] >= CONFIDENCE_THRESHOLD
        for item in preview["mapping"]
        if item["target"]
    )
    # 全角括号那一列必须由规范化层（而非别名表）命中。
    normalized = [item for item in preview["mapping"] if item["matched_by"] == "normalized"]
    assert [item["column"] for item in normalized] == ["课节时长（小时）"]

    assert preview["unmatched_columns"] == ["备注"]
    assert preview["missing_fields"] == []
    assert preview["stats"]["rows_total"] == len(_sample_rows()) + 1
    assert preview["stats"]["rows_valid"] == len(_sample_rows())
    assert preview["stats"]["rows_skipped"] == 1
    assert preview["issues"] == [
        {"行号": len(_sample_rows()) + 3, "原因": "缺少上课日期或上课时段"}
    ]


def test_preview_reruns_with_user_corrected_mapping(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    grid = _renamed_grid()
    payload = _grid_to_xlsx(grid)
    preview = _preview(client, auth_headers, payload)

    # 用户把「备注」也指定为不导入，并显式回传全部列——Fix 循环重跑。
    response = client.post(
        "/api/v1/imports/preview",
        headers=auth_headers,
        files={"file": ("renamed.xlsx", payload, "application/octet-stream")},
        data={"mapping_json": _mapping_input(preview, skip_note=False)},
    )
    assert response.status_code == 200, response.text
    rerun = response.json()
    manual = [item for item in rerun["mapping"] if item["target"] is not None]
    assert len(manual) == len(HEADERS)
    assert all(item["matched_by"] == "manual" for item in manual)
    # 没被映射覆盖的「备注」列保持未匹配，不会悄悄进入导入。
    assert rerun["stats"]["mapped_columns"] == len(HEADERS)
    assert "备注" in rerun["unmatched_columns"]
    assert rerun["issues"] == preview["issues"]


def test_preview_reports_row_level_issues_for_broken_rows(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    grid = _renamed_grid()
    preview = _preview(client, auth_headers, _grid_to_xlsx(grid))
    reasons = {issue["原因"] for issue in preview["issues"]}
    assert any("缺少上课日期" in reason for reason in reasons)


def test_commit_upserts_renamed_workbook_and_second_import_updates(
    client: TestClient, auth_headers: dict[str, str], import_set_headers: dict[str, str]
) -> None:
    payload = _grid_to_xlsx(_renamed_grid())
    preview = _preview(client, import_set_headers, payload)
    mapping_json = _mapping_input(preview)

    def commit() -> dict[str, Any]:
        response = client.post(
            "/api/v1/imports/commit",
            headers=import_set_headers,
            files={"file": ("renamed.xlsx", payload, "application/octet-stream")},
            data={"mapping_json": mapping_json, "mode": "upsert"},
        )
        assert response.status_code == 200, response.text
        return response.json()

    first = commit()
    assert first["mode"] == "upsert"
    assert first["course_sessions"] == 2
    # 与模板直通导入同构：rows_total 只计进入管线的行（坏行已在校验层跳过）。
    assert first["rows_total"] == len(_sample_rows())
    assert first["rows_skipped"] == 1
    assert first["rows_dropped_placeholder_room"] == 1
    assert first["schedule_version_no"] == 1
    assert first["orphans_deleted"] == 0

    sessions = client.get("/api/v1/course-sessions", headers=import_set_headers)
    assert sessions.status_code == 200
    assert len(sessions.json()) == 2

    second = commit()
    # 同一份表重复导入：upsert 语义下不新增课次、不产生重复。
    assert second["course_sessions"] == 0
    assert second["course_sessions_updated"] == 2
    assert second["schedule_version_no"] == 2
    sessions = client.get("/api/v1/course-sessions", headers=import_set_headers)
    assert len(sessions.json()) == 2


def test_commit_insert_mode_keeps_existing_rows_untouched(
    client: TestClient, auth_headers: dict[str, str], import_set_headers: dict[str, str]
) -> None:
    payload = _grid_to_xlsx(_renamed_grid())
    preview = _preview(client, import_set_headers, payload)

    def commit() -> dict[str, Any]:
        response = client.post(
            "/api/v1/imports/commit",
            headers=import_set_headers,
            files={"file": ("renamed.xlsx", payload, "application/octet-stream")},
            data={"mapping_json": _mapping_input(preview), "mode": "insert"},
        )
        assert response.status_code == 200, response.text
        return response.json()

    first = commit()
    assert first["mode"] == "insert"
    assert first["course_sessions"] == 2

    second = commit()
    assert second["course_sessions"] == 0
    assert second["course_sessions_skipped_existing"] == 2
    sessions = client.get("/api/v1/course-sessions", headers=import_set_headers)
    assert len(sessions.json()) == 2


def test_commit_rejects_invalid_mapping(client: TestClient, auth_headers: dict[str, str]) -> None:
    payload = _grid_to_xlsx(_renamed_grid())

    missing = client.post(
        "/api/v1/imports/commit",
        headers=auth_headers,
        files={"file": ("renamed.xlsx", payload, "application/octet-stream")},
    )
    assert missing.status_code == 422

    unknown_column = client.post(
        "/api/v1/imports/commit",
        headers=auth_headers,
        files={"file": ("renamed.xlsx", payload, "application/octet-stream")},
        data={
            "mapping_json": json.dumps(
                {"columns": [{"column": "不存在的列", "target": "上课日期"}]},
                ensure_ascii=False,
            )
        },
    )
    assert unknown_column.status_code == 422
    assert "不存在的列" in unknown_column.json()["detail"]

    duplicate_target = client.post(
        "/api/v1/imports/commit",
        headers=auth_headers,
        files={"file": ("renamed.xlsx", payload, "application/octet-stream")},
        data={
            "mapping_json": json.dumps(
                {
                    "columns": [
                        {"column": "日期", "target": "上课日期"},
                        {"column": "任课教师", "target": "上课日期"},
                    ]
                },
                ensure_ascii=False,
            )
        },
    )
    assert duplicate_target.status_code == 422
    assert "只能对应一列" in duplicate_target.json()["detail"]


def test_preview_rejects_unsupported_files(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/imports/preview",
        headers=auth_headers,
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert response.status_code == 400
    assert "只接受" in response.json()["detail"]
