"""郑州官方 14 列模板的直通导入入口。

「从规范数据到入库」的完整管线已抽取到 :mod:`app.services.converter_core`，
本模块只负责「工作簿文件 → 14 列记录行」的解析与表头校验，对外行为与旧版
完全一致；L2 智能映射导入也驱动同一条核心管线。
"""

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..db import SessionLocal
from .converter_core import (  # noqa: F401  再导出：既有调用方仍从本模块取这些名字
    CAMPUS_BUSINESS_ID,
    CAMPUS_NAME,
    OFFICIAL_VERSION_SUFFIX,
    PLACEHOLDER_ROOM,
    TEMPLATE_HEADERS,
    WEEKDAYS,
    WorkbookFormatError,
    _collect_lesson_rows,
    _row_identity,
    _split_placeholder_rows,
    import_canonical_rows,
    official_version_name,
    parse_template_records,
)

SHEET_NAME = "课表数据源"


def _validate_headers(sheet: Any) -> None:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    header = tuple(
        (str(value).strip() if value is not None else "") for value in (header_row or ())
    )
    if header[: len(TEMPLATE_HEADERS)] == TEMPLATE_HEADERS:
        return
    mismatches: list[str] = []
    for index, expected in enumerate(TEMPLATE_HEADERS):
        actual = header[index] if index < len(header) else ""
        if actual != expected:
            mismatches.append(f"第 {index + 1} 列应为「{expected}」，实际为「{actual}」")
    raise WorkbookFormatError(
        "工作簿列名与官方模板不一致，请下载模板后按列填写：" + "；".join(mismatches[:5])
    )


def _read_template_records(workbook_path: Path) -> list[dict[str, Any]]:
    """读取官方模板工作簿，返回「规范表头 → 单元格值」的记录行（含空行）。"""
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise WorkbookFormatError(
            f"工作簿缺少「{SHEET_NAME}」工作表，实际为：{workbook.sheetnames}"
        )
    sheet = workbook[SHEET_NAME]
    _validate_headers(sheet)
    records: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        records.append(
            {
                header: (values[index] if index < len(values) else None)
                for index, header in enumerate(TEMPLATE_HEADERS)
            }
        )
    workbook.close()
    return records


def _read_rows(workbook_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """读取官方模板工作簿，返回（内部规范行, 跳过行清单）。"""
    rows, skipped, _blank_rows = parse_template_records(_read_template_records(workbook_path))
    return rows, skipped


def import_schedule_workbook(
    db: Session,
    workbook_path: Path,
    campus_business_id: str = CAMPUS_BUSINESS_ID,
    campus_name: str = CAMPUS_NAME,
    schedule_set_id: str = "default",
) -> dict[str, Any]:
    """官方模板直通导入：14 列表头严格匹配后走统一核心管线。"""
    return import_canonical_rows(
        db,
        _read_template_records(workbook_path),
        campus_business_id=campus_business_id,
        campus_name=campus_name,
        schedule_set_id=schedule_set_id,
        source_name=workbook_path.name,
        checksum=hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
    )


def import_zhengzhou(db: Session, workbook_path: Path) -> dict[str, Any]:
    return import_schedule_workbook(db, workbook_path)


def main() -> None:
    if len(sys.argv) < 2:
        print("用法：python -m app.services.converter_zhengzhou <xlsx 路径>")
        raise SystemExit(2)
    workbook_path = Path(sys.argv[1])
    with SessionLocal() as db:
        result = import_schedule_workbook(db, workbook_path)
        db.commit()
        print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
