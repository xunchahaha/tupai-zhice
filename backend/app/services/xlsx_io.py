from __future__ import annotations

import json
from io import BytesIO
from typing import cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Room, ScheduleAssignment, ScheduleVersion, Teacher, TimeSlot
from .snapshot import version_course_map


def export_schedule_xlsx(db: Session, schedule: ScheduleVersion) -> bytes:
    assignments = list(
        db.scalars(
            select(ScheduleAssignment)
            .where(ScheduleAssignment.schedule_version_id == schedule.id)
            .order_by(
                ScheduleAssignment.lesson_date,
                ScheduleAssignment.room_business_id,
                ScheduleAssignment.course_session_id,
            )
        )
    )
    courses = version_course_map(db, schedule)
    rooms = {
        item.business_id: item
        for item in db.scalars(
            select(Room).where(Room.schedule_set_id == schedule.schedule_set_id)
        )
    }
    teachers = {
        item.business_id: item
        for item in db.scalars(
            select(Teacher).where(Teacher.schedule_set_id == schedule.schedule_set_id)
        )
    }
    slots = {
        item.business_id: item
        for item in db.scalars(
            select(TimeSlot).where(TimeSlot.schedule_set_id == schedule.schedule_set_id)
        )
    }
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "课表"
    headers = [
        "版本",
        "场次ID",
        "业务线",
        "产品班型",
        "班级ID",
        "教师ID",
        "具体日程账号",
        "课节名称",
        "编排阶段",
        "上课日期",
        "实际开始时间",
        "实际结束时间",
        "固定开始时间",
        "固定结束时间",
        "时段ID",
        "原始教室ID",
        "原始教室名称",
        "最终教室ID",
        "最终教室名称",
        "变更状态",
    ]
    sheet.append(headers)
    for assignment in assignments:
        course = courses[assignment.course_session_id]
        original_room = rooms.get(course.original_room_business_id or "")
        final_room = rooms.get(assignment.room_business_id)
        teacher = teachers.get(course.teacher_business_id)
        slot = slots.get(assignment.slot_business_id)
        calendar_user_id = course.calendar_user_id or (
            teacher.calendar_user_id if teacher else None
        )
        sheet.append(
            [
                schedule.version_no,
                course.business_id,
                course.business_line,
                " / ".join(course.product_types or [course.product_type]),
                course.class_business_id,
                course.teacher_business_id,
                calendar_user_id or "",
                " / ".join(course.lesson_names or [course.lesson_name]),
                " / ".join(course.stages or [course.stage]),
                assignment.lesson_date.isoformat() if assignment.lesson_date else "",
                slot.start_time if slot else course.fixed_start_time,
                slot.end_time if slot else course.fixed_end_time,
                course.fixed_start_time,
                course.fixed_end_time,
                assignment.slot_business_id,
                course.original_room_business_id or "",
                original_room.name if original_room else "",
                assignment.room_business_id,
                final_room.name if final_room else "",
                assignment.change_kind,
            ]
        )
    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = [
        10,
        16,
        12,
        32,
        16,
        16,
        22,
        32,
        22,
        12,
        14,
        14,
        14,
        14,
        18,
        16,
        20,
        16,
        20,
        12,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    metrics = workbook.create_sheet("指标")
    metrics.append(["指标", "值"])
    for key, value in schedule.metrics.items():
        # 指标里有分维度明细这类结构化值，Excel 单元格只接受标量。
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        metrics.append([key, value])
    metrics.column_dimensions["A"].width = 28
    metrics.column_dimensions["B"].width = 18
    for cell in metrics[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
