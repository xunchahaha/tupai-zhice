from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Campus, ClassGroup, CourseSession, Room, Rule, Teacher, TimeSlot, User
from ..schemas import ImportResult
from ..security import hash_password


def split_values(value: Any) -> list[str]:
    if value is None or str(value).strip() == "":
        return []
    return [item.strip() for item in str(value).split("|") if item.strip()]


def as_bool(value: Any) -> bool:
    return str(value).strip().lower() in {"是", "true", "1", "yes"}


def as_int(value: Any, default: int = 0) -> int:
    if value is None or str(value).strip() == "":
        return default
    return int(float(str(value)))


def bootstrap_admin(db: Session, username: str, password: str) -> User:
    user = db.scalar(select(User).where(User.username == username))
    if user:
        return user
    user = User(username=username, password_hash=hash_password(password), role="admin")
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def _upsert(db: Session, model: type[Any], match: dict[str, Any], values: dict[str, Any]) -> Any:
    statement = select(model)
    for key, value in match.items():
        statement = statement.where(getattr(model, key) == value)
    instance = db.scalar(statement)
    if instance is None:
        instance = model(**match, **values)
        db.add(instance)
    else:
        for key, value in values.items():
            setattr(instance, key, value)
    return instance


def import_sample_workbook(db: Session, workbook_path: Path) -> ImportResult:
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    campus = db.scalar(select(Campus).where(Campus.business_id == "CAMPUS-DEMO"))
    if campus is None:
        campus = Campus(business_id="CAMPUS-DEMO", name="示范校区")
        db.add(campus)
        db.flush()

    counts = {
        "teachers": 0,
        "class_groups": 0,
        "rooms": 0,
        "time_slots": 0,
        "course_sessions": 0,
        "rules": 0,
    }

    for row in workbook["教师"].iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        _upsert(
            db,
            Teacher,
            {"campus_id": campus.id, "business_id": str(row[0])},
            {
                "name": str(row[1]),
                "subject": str(row[2] or ""),
                "max_hours": as_int(row[3], 8),
                "unavailable_slot_ids": split_values(row[4]),
                "preferred_slot_ids": split_values(row[5]),
                "data_level": str(row[6] or "内部"),
            },
        )
        counts["teachers"] += 1

    for row in workbook["班级"].iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        _upsert(
            db,
            ClassGroup,
            {"campus_id": campus.id, "business_id": str(row[0])},
            {
                "name": str(row[1]),
                "grade": str(row[2] or ""),
                "subject": str(row[3] or ""),
                "student_count": as_int(row[4]),
                "priority": str(row[5] or "常规"),
                "required_devices": split_values(row[6]),
                "teacher_business_id": str(row[7]),
            },
        )
        counts["class_groups"] += 1

    for row in workbook["教室"].iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        _upsert(
            db,
            Room,
            {"campus_id": campus.id, "business_id": str(row[0])},
            {
                "name": str(row[1]),
                "capacity": as_int(row[2]),
                "devices": split_values(row[3]),
                "available_slot_ids": split_values(row[4]),
                "is_active": True,
            },
        )
        counts["rooms"] += 1

    for sequence, row in enumerate(
        workbook["时段"].iter_rows(min_row=5, values_only=True), start=1
    ):
        if not row[0]:
            continue
        _upsert(
            db,
            TimeSlot,
            {"campus_id": campus.id, "business_id": str(row[0])},
            {
                "weekday": str(row[1]),
                "start_time": str(row[2]),
                "end_time": str(row[3]),
                "kind": str(row[4] or ""),
                "sequence": sequence,
                "is_open": as_bool(row[5]),
            },
        )
        counts["time_slots"] += 1

    for row in workbook["课程需求"].iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        _upsert(
            db,
            CourseSession,
            {"campus_id": campus.id, "business_id": str(row[0])},
            {
                "class_business_id": str(row[1]),
                "teacher_business_id": str(row[2]),
                "subject": str(row[3] or ""),
                "student_count": as_int(row[4]),
                "required_devices": split_values(row[5]),
                "duration_minutes": as_int(row[6], 90),
                "suggested_slot_id": str(row[7]) if row[7] else None,
                "is_locked": False,
            },
        )
        counts["course_sessions"] += 1

    for row in workbook["规则样本"].iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        hardness = "hard" if str(row[4]) == "硬" else "soft"
        _upsert(
            db,
            Rule,
            {"business_id": str(row[0])},
            {
                "source_text": str(row[1]),
                "actor_type": str(row[2] or "system"),
                "actor_ids": [],
                "constraint_type": str(row[3] or "declared_constraint"),
                "scope": {},
                "hardness": hardness,
                "weight": None if hardness == "hard" else as_int(row[6], 1),
                "structured_expression": {"expression": str(row[5] or "")},
                "source_doc": "数据分析样本/规则样本",
                "confidence": 1.0,
                "status": "active" if str(row[7]) == "已确认" else "awaiting_confirmation",
            },
        )
        counts["rules"] += 1

    db.commit()
    return ImportResult(
        source=str(workbook_path),
        campuses=1,
        **counts,
    )
